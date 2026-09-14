from __future__ import annotations

import io
from typing import Any

from openpyxl import load_workbook

from src.parsers.base import ParseError
from src.parsers.mapping import resolve_row


class ExcelParser:
    """อ่านเฉพาะ sheet แรกที่ active — ถ้าต้องรองรับหลาย sheet ค่อยขยายทีหลัง"""

    def parse(self, raw: bytes, mapping_name: str = "default") -> list[dict[str, Any]]:
        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as exc:
            raise ParseError(f"เปิดไฟล์ Excel ไม่ได้: {exc}") from exc

        sheet = wb.active
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration as exc:
            raise ParseError("ไฟล์ Excel ไม่มีข้อมูล") from exc

        headers = [str(h) if h is not None else "" for h in header]

        results: list[dict[str, Any]] = []
        for i, values in enumerate(rows_iter):
            if all(v is None for v in values):
                continue  # ข้าม row ว่างที่ Excel มักทิ้งไว้ท้ายไฟล์
            row = dict(zip(headers, values, strict=False))
            try:
                results.append(resolve_row(row, mapping_name))
            except Exception as exc:
                raise ParseError(f"Excel row {i + 2}: {exc}") from exc
        return results
